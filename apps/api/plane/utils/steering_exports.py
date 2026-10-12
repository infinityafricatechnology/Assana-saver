# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.

"""
Exports of the steering space (cahier des charges §6.7, §12): the COMEX PDF, A4 landscape, and the
Excel file of a filtered list. Both are French documents marked « Confidentiel », like the weekly
report built on the PDF (§9.3).
"""

# Python imports
import io
import os
from datetime import date, datetime
from typing import Any, Dict, Iterable, List, Optional
from xml.sax.saxutils import escape

# Third party imports
import reportlab
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

# Module imports
from plane.db.models import SteeringRiskNature, SteeringStatus, SteeringWaitingFor
from plane.utils.steering_dashboard import ABIDJAN, PRIORITY_RANK, Portfolio

# §11, the Group's charter
GOLD = colors.HexColor("#E8B93E")
GOLD_TEXT = colors.HexColor("#9A7622")
CREAM = colors.HexColor("#F3F0E9")
BORDER = colors.HexColor("#E8E3D8")
TEXT = colors.HexColor("#1B1B19")
TEXT_2 = colors.HexColor("#5E5B53")
BANNER = colors.HexColor("#1D1D1B")
SIGNAL_COLORS = {
    "critical": colors.HexColor("#B23B2A"),
    "late": colors.HexColor("#C0701A"),
    "soon": colors.HexColor("#2B4A80"),
    "on_time": colors.HexColor("#2E6B52"),
    "closed": TEXT_2,
    "queue": colors.HexColor("#6A4C9E"),
}
STATE_COLORS = {"alert": SIGNAL_COLORS["critical"], "watch": SIGNAL_COLORS["late"], "ok": SIGNAL_COLORS["on_time"]}

PRIORITY_LABELS = {"urgent": "Urgent", "high": "Important", "medium": "Normal"}
SIGNAL_LABELS = {
    "critical": "Critique",
    "late": "En retard",
    "soon": "Échéance proche",
    "on_time": "Dans les délais",
    "closed": "Clôturée",
}
STATE_LABELS = {"alert": "En alerte", "watch": "À surveiller", "ok": "Sans alerte"}
ROLE_LABELS = {"ceo": "le PDG", "cabinet": "le Cabinet"}
STATUS_LABELS = dict(SteeringStatus.choices)
WAITING_LABELS = dict(SteeringWaitingFor.choices)
RISK_LABELS = dict(SteeringRiskNature.choices)
MONTHS = ("janv.", "févr.", "mars", "avr.", "mai", "juin", "juil.", "août", "sept.", "oct.", "nov.", "déc.")

FONT, FONT_BOLD = "IAG", "IAG-Bold"
_fonts_registered = False


def _register_fonts() -> None:
    """Bitstream Vera, shipped with reportlab, covers the French accents and the punctuation used here."""
    global _fonts_registered
    if _fonts_registered:
        return
    directory = os.path.join(os.path.dirname(reportlab.__file__), "fonts")
    pdfmetrics.registerFont(TTFont(FONT, os.path.join(directory, "Vera.ttf")))
    pdfmetrics.registerFont(TTFont(FONT_BOLD, os.path.join(directory, "VeraBd.ttf")))
    pdfmetrics.registerFontFamily(FONT, normal=FONT, bold=FONT_BOLD, italic=FONT, boldItalic=FONT_BOLD)
    _fonts_registered = True


# ---------------------------------------------------------------------------------------------------
# French formats (§12: « 7 oct. 2026 »)
# ---------------------------------------------------------------------------------------------------


def french_date(value: Any) -> str:
    if not value:
        return ""
    if isinstance(value, str):
        value = date.fromisoformat(value[:10])
    if isinstance(value, datetime):
        value = value.astimezone(ABIDJAN).date()
    return f"{value.day} {MONTHS[value.month - 1]} {value.year}"


def delay_label(row: Dict[str, Any]) -> str:
    if not row["is_open"]:
        return ""
    if row["days_late"] > 0:
        return f"J+{row['days_late']}"
    if row["days_left"] == 0:
        return "Aujourd'hui"
    if row["days_left"] is not None:
        return f"Dans {row['days_left']} j"
    return ""


def risk_label(row: Dict[str, Any]) -> str:
    risk = row.get("risk")
    if not risk:
        return ""
    parts = [RISK_LABELS.get(risk["nature"], "") if risk["nature"] else ""]
    if risk["effective_date"]:
        days = risk["days"]
        countdown = f"J-{days}" if days is not None and days >= 0 else f"J+{-days}" if days is not None else ""
        parts.append(f"{french_date(risk['effective_date'])} ({countdown})")
    return " · ".join(part for part in parts if part)


# ---------------------------------------------------------------------------------------------------
# COMEX PDF (§6.7)
# ---------------------------------------------------------------------------------------------------


def _styles() -> Dict[str, ParagraphStyle]:
    base = ParagraphStyle("base", fontName=FONT, fontSize=8, leading=10, textColor=TEXT)
    return {
        "cell": base,
        "cell_muted": ParagraphStyle("muted", parent=base, textColor=TEXT_2),
        "cell_bold": ParagraphStyle("bold", parent=base, fontName=FONT_BOLD),
        "header": ParagraphStyle("header", parent=base, fontName=FONT_BOLD, textColor=TEXT_2, fontSize=7.5),
        "title": ParagraphStyle(
            "title",
            parent=base,
            fontName=FONT_BOLD,
            fontSize=12,
            leading=15,
            spaceBefore=8,
            spaceAfter=4,
            keepWithNext=1,
        ),
        "kpi_label": ParagraphStyle("kpi_label", parent=base, textColor=TEXT_2, fontSize=7.5),
        "empty": ParagraphStyle("empty", parent=base, textColor=TEXT_2),
    }


def _p(text: Any, style: ParagraphStyle, color: Optional[colors.Color] = None) -> Paragraph:
    content = escape("" if text is None else str(text))
    if color is not None:
        content = f'<font color="{color.hexval()}">{content}</font>'
    return Paragraph(content, style)


def _table(
    header: List[str], rows: List[List[Any]], widths: List[float], styles, edges: Optional[List] = None
) -> Table:
    data = [[_p(label, styles["header"]) for label in header], *rows]
    table = Table(data, colWidths=widths, repeatRows=1)
    commands = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, BORDER),
        ("BACKGROUND", (0, 0), (-1, 0), CREAM),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]
    # A signal reads on the left of its line too (§11).
    for position, color in enumerate(edges or [], start=1):
        if color is not None:
            commands.append(("LINEBEFORE", (0, position), (0, position), 2.5, color))
    table.setStyle(TableStyle(commands))
    return table


class _NumberedCanvas(canvas.Canvas):
    """Draws « Confidentiel · Page n / N » once the number of pages is known."""

    footer_text = ""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._pages: List[Dict[str, Any]] = []

    def showPage(self):
        self._pages.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._pages)
        for state in self._pages:
            self.__dict__.update(state)
            width, _ = self._pagesize
            self.setFont(FONT, 7.5)
            self.setFillColor(TEXT_2)
            self.drawString(12 * mm, 8 * mm, self.footer_text)
            self.drawRightString(width - 12 * mm, 8 * mm, f"Confidentiel · Page {self._pageNumber} / {total}")
            super().showPage()
        super().save()


def comex_pdf(portfolio: Portfolio, branch_id: Optional[str] = None) -> bytes:
    """
    The COMEX export: banner and situation date, the six indicators, the situation by entity (the
    ten most late at least, with their branch), the heaviest delays, the CEO queue in the order of
    the tab with priority and risk, and the blocking work items in appendix. It follows the
    branch filter of the screen.
    """
    _register_fonts()
    styles = _styles()
    summary = portfolio.summary(branch_id)
    rows = {row["id"]: row for row in portfolio.rows}
    branch_name = portfolio.branches.get(branch_id, {}).get("name") if branch_id else None
    situation = (
        f"Situation au {french_date(portfolio.today)} à "
        f"{portfolio.generated_at.astimezone(ABIDJAN).strftime('%H:%M')}"
        + (f" · Pôle {branch_name}" if branch_name else " · Tout le Groupe")
    )
    page_width, page_height = landscape(A4)
    content_width = page_width - 24 * mm
    entity_name = lambda entity_id: portfolio.entities.get(entity_id, {}).get("name", "")  # noqa: E731
    people = lambda ids: ", ".join(portfolio.people.get(i, {}).get("name", "") for i in ids)  # noqa: E731

    story: List[Any] = []

    # Six indicators (§6.1)
    kpis = summary["kpis"]
    tiles = [
        ("Tâches ouvertes", kpis["open"], TEXT, f"Avancement moyen {kpis['average_progress']} %"),
        ("En retard", kpis["late"], SIGNAL_COLORS["late"], f"{kpis['late_share']} % des ouvertes"),
        ("Critiques > 30 j", kpis["critical"], SIGNAL_COLORS["critical"], ""),
        ("Bloquées", kpis["blocked"], SIGNAL_COLORS["critical"], ""),
        ("File PDG", kpis["queue"], SIGNAL_COLORS["queue"], ""),
        ("Échéances à 14 jours", kpis["due_soon"], SIGNAL_COLORS["soon"], ""),
    ]
    value_style = ParagraphStyle("kpi_value", parent=styles["cell_bold"], fontSize=20, leading=24)
    kpi_cells = [
        [_p(label, styles["kpi_label"]), _p(value, value_style, color), _p(detail, styles["kpi_label"])]
        for label, value, color, detail in tiles
    ]
    kpi_table = Table([kpi_cells], colWidths=[content_width / 6] * 6)
    kpi_table.setStyle(
        TableStyle(
            [
                ("BOX", (0, 0), (-1, -1), 0.5, BORDER),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, BORDER),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                *[("LINEBEFORE", (i, 0), (i, 0), 2.5, tile[2]) for i, tile in enumerate(tiles) if tile[2] is not TEXT],
            ]
        )
    )
    story += [kpi_table, Spacer(1, 6)]

    # Situation by entity: the ten most late at least (§6.7)
    entities = [{**entity, "branch": branch["name"]} for branch in summary["branches"] for entity in branch["entities"]]
    entities.sort(key=lambda entity: (-entity["late"], -entity["critical"], entity["name"]))
    entity_rows = [
        [
            _p(entity["name"], styles["cell_bold"]),
            _p(entity["branch"], styles["cell_muted"]),
            _p(entity["open"], styles["cell"]),
            _p(entity["late"], styles["cell"], SIGNAL_COLORS["late"] if entity["late"] else None),
            _p(entity["critical"], styles["cell"], SIGNAL_COLORS["critical"] if entity["critical"] else None),
            _p(entity["blocked"], styles["cell"], SIGNAL_COLORS["critical"] if entity["blocked"] else None),
            _p(entity["queue"], styles["cell"], SIGNAL_COLORS["queue"] if entity["queue"] else None),
            _p(STATE_LABELS[entity["state"]], styles["cell"], STATE_COLORS[entity["state"]]),
        ]
        for entity in entities[: max(10, sum(entity["late"] > 0 for entity in entities))]
    ]
    story.append(_p("Situation par entité", styles["title"]))
    if entity_rows:
        widths = [0.26, 0.2, 0.08, 0.08, 0.08, 0.08, 0.08, 0.14]
        story.append(
            _table(
                ["Entité", "Pôle", "Ouvertes", "En retard", "Critiques", "Bloquées", "File PDG", "État"],
                entity_rows,
                [content_width * width for width in widths],
                styles,
                [STATE_COLORS[entity["state"]] for entity in entities[: len(entity_rows)]],
            )
        )
    else:
        story.append(_p("Aucune entité dans le périmètre.", styles["empty"]))

    # Heaviest delays, in the order of RG-12
    delays = sorted(
        (row for row in portfolio.rows_for(branch_id) if row["is_open"] and row["days_late"] > 0),
        key=lambda row: (PRIORITY_RANK.get(row["priority"], 2), -row["days_late"]),
    )[:15]
    story.append(_p("Retards les plus lourds", styles["title"]))
    if delays:
        story.append(
            _table(
                ["Clé", "Tâche", "Entité", "Responsable", "Priorité", "Échéance", "Retard"],
                [
                    [
                        _p(row["key"], styles["cell_muted"]),
                        _p(row["name"], styles["cell_bold"]),
                        _p(entity_name(row["entity_id"]), styles["cell"]),
                        _p(people(row["assignee_ids"]), styles["cell"]),
                        _p(PRIORITY_LABELS.get(row["priority"], ""), styles["cell"]),
                        _p(french_date(row["target_date"]), styles["cell"]),
                        _p(delay_label(row), styles["cell_bold"], SIGNAL_COLORS[row["signal"]]),
                    ]
                    for row in delays
                ],
                [content_width * width for width in (0.08, 0.34, 0.18, 0.16, 0.08, 0.1, 0.06)],
                styles,
                [SIGNAL_COLORS[row["signal"]] for row in delays],
            )
        )
    else:
        story.append(_p("Aucune tâche en retard.", styles["empty"]))

    # CEO queue, in the order of the tab (RG-04)
    queue = [row for row in portfolio.queue() if not branch_id or row["branch_id"] == branch_id]
    story.append(_p("File PDG", styles["title"]))
    if queue:
        story.append(
            _table(
                ["Priorité retenue", "Clé", "Tâche", "Entité", "Échéance", "Attente", "Risque si non validée"],
                [
                    [
                        _p(
                            PRIORITY_LABELS.get(row["retained_priority"], "")
                            + (
                                f" (requalifiée par {ROLE_LABELS.get(row['requalification']['by_role'], '')})"
                                if row["requalification"]
                                else ""
                            ),
                            styles["cell_bold"],
                        ),
                        _p(row["key"], styles["cell_muted"]),
                        _p(row["name"], styles["cell_bold"]),
                        _p(entity_name(row["entity_id"]), styles["cell"]),
                        _p(f"{french_date(row['target_date'])} {delay_label(row)}".strip(), styles["cell"]),
                        _p(f"{row['submission']['waiting_days']} j", styles["cell"]),
                        _p(
                            risk_label(row)
                            + (
                                f" — {row['risk']['description']}" if row["risk"] and row["risk"]["description"] else ""
                            ),
                            styles["cell"],
                            SIGNAL_COLORS["critical"] if row["risk"] and row["risk"]["soon"] else None,
                        ),
                    ]
                    for row in queue
                ],
                [content_width * width for width in (0.13, 0.07, 0.26, 0.14, 0.12, 0.06, 0.22)],
                styles,
                [SIGNAL_COLORS["queue"]] * len(queue),
            )
        )
    else:
        story.append(_p("Aucune tâche n'attend la décision du PDG.", styles["empty"]))

    # Appendix: blocking work items (§7)
    blocking = [(rows[item["id"]], item) for item in summary["blocking"]]
    story.append(_p("Annexe — Tâches bloquantes", styles["title"]))
    if blocking:
        key_of = lambda item_id: (rows.get(item_id) or portfolio.related.get(item_id) or {}).get("key", "")  # noqa: E731
        story.append(
            _table(
                ["Clé", "Tâche bloquante", "Entité", "État", "Bloque", "Tâches bloquées"],
                [
                    [
                        _p(row["key"], styles["cell_muted"]),
                        _p(row["name"], styles["cell_bold"]),
                        _p(entity_name(row["entity_id"]), styles["cell"]),
                        _p(
                            delay_label(row) if row["days_late"] > 0 else STATUS_LABELS.get(row["status"], ""),
                            styles["cell"],
                            SIGNAL_COLORS[row["signal"]] if row["days_late"] > 0 else None,
                        ),
                        _p(item["count"], styles["cell"]),
                        _p(", ".join(key_of(blocked) for blocked in row["blocking"]), styles["cell"]),
                    ]
                    for row, item in blocking
                ],
                [content_width * width for width in (0.08, 0.34, 0.18, 0.12, 0.06, 0.22)],
                styles,
                [SIGNAL_COLORS[row["signal"]] if row["days_late"] > 0 else None for row, _ in blocking],
            )
        )
    else:
        story.append(_p("Aucune tâche n'en bloque une autre.", styles["empty"]))

    def first_page(pdf: canvas.Canvas, _doc) -> None:
        # The banner: anthracite, rounded at the bottom, with its gold line (§11).
        banner_height = 26 * mm
        pdf.saveState()
        pdf.setFillColor(BANNER)
        pdf.roundRect(0, page_height - banner_height, page_width, banner_height + 10 * mm, 8 * mm, stroke=0, fill=1)
        pdf.rect(0, page_height - banner_height + 8 * mm, page_width, banner_height, stroke=0, fill=1)
        pdf.setStrokeColor(GOLD)
        pdf.setLineWidth(1.2)
        pdf.line(
            12 * mm,
            page_height - banner_height + 1.5 * mm,
            page_width - 12 * mm,
            page_height - banner_height + 1.5 * mm,
        )
        pdf.setFillColor(GOLD)
        pdf.setFont(FONT_BOLD, 7.5)
        pdf.drawString(12 * mm, page_height - 8 * mm, "INFINITY AFRICA GROUP")
        pdf.setFillColor(colors.white)
        pdf.setFont(FONT_BOLD, 15)
        pdf.drawString(12 * mm, page_height - 15 * mm, "Espace de pilotage — Rapport COMEX")
        pdf.setFillColor(colors.HexColor("#D9D3C4"))
        pdf.setFont(FONT, 8.5)
        pdf.drawString(12 * mm, page_height - 21 * mm, situation)
        pdf.setFillColor(GOLD)
        pdf.drawRightString(page_width - 12 * mm, page_height - 8 * mm, "Confidentiel")
        pdf.restoreState()

    def later_pages(pdf: canvas.Canvas, _doc) -> None:
        pdf.saveState()
        pdf.setFont(FONT_BOLD, 8)
        pdf.setFillColor(TEXT)
        pdf.drawString(12 * mm, page_height - 10 * mm, "Espace de pilotage — Rapport COMEX")
        pdf.setFont(FONT, 8)
        pdf.setFillColor(TEXT_2)
        pdf.drawRightString(page_width - 12 * mm, page_height - 10 * mm, situation)
        pdf.setStrokeColor(GOLD)
        pdf.line(12 * mm, page_height - 12.5 * mm, page_width - 12 * mm, page_height - 12.5 * mm)
        pdf.restoreState()

    # The first page leaves room for the banner; the others only for a slim header line.
    margin, bottom = 12 * mm, 15 * mm
    frames = {
        "first": Frame(margin, bottom, content_width, page_height - bottom - 32 * mm, 0, 0, 0, 0, id="first"),
        "later": Frame(margin, bottom, content_width, page_height - bottom - 18 * mm, 0, 0, 0, 0, id="later"),
    }
    buffer = io.BytesIO()
    document = BaseDocTemplate(
        buffer,
        pagesize=(page_width, page_height),
        title="Espace de pilotage — Rapport COMEX",
        author="Infinity Planning",
        subject=situation,
    )
    document.addPageTemplates(
        [
            PageTemplate(id="first", frames=[frames["first"]], onPage=first_page, autoNextPageTemplate="later"),
            PageTemplate(id="later", frames=[frames["later"]], onPage=later_pages),
        ]
    )
    footer = f"Infinity Planning · {situation}"

    class Footer(_NumberedCanvas):
        footer_text = footer

    document.build([KeepTogether(story[:2]), *story[2:]], canvasmaker=Footer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------------------------------
# Excel export of a filtered list (§12 « Export de données »)
# ---------------------------------------------------------------------------------------------------

XLSX_COLUMNS = (
    ("Clé", 10),
    ("Tâche", 48),
    ("Projet", 24),
    ("Pôle", 22),
    ("Entité", 26),
    ("Catégorie", 18),
    ("Responsable", 24),
    ("Superviseur", 22),
    ("Approbateur", 22),
    ("Priorité déclarée", 12),
    ("Priorité retenue", 12),
    ("Échéance", 12),
    ("Retard (j)", 9),
    ("Signal", 16),
    ("Statut", 14),
    ("Avancement (%)", 11),
    ("En attente de", 18),
    ("En attente depuis", 14),
    ("File PDG", 9),
    ("Risque", 16),
    ("Date d'effet du risque", 14),
    ("Bloquée par", 22),
    ("Bloque", 22),
    ("Incohérences", 28),
    ("Point de situation", 48),
    ("Mise à jour", 14),
)

INCONSISTENCY_LABELS = {
    "progress_without_closure": "Avancement à 100 % sans clôture",
    "late_not_started": "Deadline dépassée et toujours « À démarrer »",
}


def tasks_xlsx(portfolio: Portfolio, ids: Optional[Iterable[str]] = None) -> bytes:
    """The work items of the list in the order of the screen; ``ids`` outside the scope are ignored."""
    by_id = {row["id"]: row for row in portfolio.rows}
    rows = [by_id[item] for item in ids if item in by_id] if ids is not None else portfolio.rows
    key_of = lambda item_id: (by_id.get(item_id) or portfolio.related.get(item_id) or {}).get("key", "")  # noqa: E731
    name = lambda person_id: portfolio.people.get(person_id, {}).get("name", "") if person_id else ""  # noqa: E731

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tâches"
    sheet.append([label for label, _ in XLSX_COLUMNS])
    for row in rows:
        entity = portfolio.entities.get(row["entity_id"], {})
        sheet.append(
            [
                row["key"],
                row["name"],
                portfolio.projects.get(row["project_id"], {}).get("name", ""),
                portfolio.branches.get(entity.get("branch_id"), {}).get("name", ""),
                entity.get("name", ""),
                portfolio.categories.get(row["category_id"], "") if row["category_id"] else "",
                ", ".join(name(person) for person in row["assignee_ids"]),
                name(row["supervisor_id"]),
                name(row["approver_id"]),
                PRIORITY_LABELS.get(row["priority"], ""),
                PRIORITY_LABELS.get(row["retained_priority"], ""),
                date.fromisoformat(row["target_date"]) if row["target_date"] else None,
                row["days_late"] or None,
                SIGNAL_LABELS.get(row["signal"], ""),
                STATUS_LABELS.get(row["status"], ""),
                row["progress"],
                WAITING_LABELS.get(row["waiting_for"], "") if row["waiting_for"] else "",
                date.fromisoformat(row["waiting_since"]) if row["waiting_since"] else None,
                "Oui" if row["in_queue"] else "",
                RISK_LABELS.get(row["risk"]["nature"], "") if row["risk"] and row["risk"]["nature"] else "",
                date.fromisoformat(row["risk"]["effective_date"])
                if row["risk"] and row["risk"]["effective_date"]
                else None,
                ", ".join(key_of(item) for item in row["blocked_by"]),
                ", ".join(key_of(item) for item in row["blocking"]),
                " ; ".join(INCONSISTENCY_LABELS[reason] for reason in row["inconsistencies"]),
                row["situation"],
                datetime.fromisoformat(row["updated_at"]).astimezone(ABIDJAN).replace(tzinfo=None)
                if row["updated_at"]
                else None,
            ]
        )

    header_fill = PatternFill("solid", fgColor="E8B93E")
    for column, (_, width) in enumerate(XLSX_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column)
        cell.font = Font(bold=True, color="1B1B19")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(column)].width = width
    for column in (12, 18, 21):
        for cell in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
            for item in cell:
                item.number_format = "DD/MM/YYYY"
    for item in sheet.iter_cols(min_col=26, max_col=26, min_row=2):
        for cell in item:
            cell.number_format = "DD/MM/YYYY HH:MM"
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions

    about = workbook.create_sheet("À propos")
    about.append(["Infinity Planning · Espace de pilotage"])
    generated = portfolio.generated_at.astimezone(ABIDJAN).strftime("%H:%M")
    about.append([f"Situation au {french_date(portfolio.today)} à {generated}"])
    about.append([f"{len(rows)} tâche(s) exportée(s)"])
    about.append(["Confidentiel"])
    about["A1"].font = Font(bold=True)
    about.column_dimensions["A"].width = 60

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
