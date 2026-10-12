# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.

"""Read a list of people from a CSV or Excel file, and write the list back as CSV."""

# Python imports
import csv
import io
import unicodedata
from typing import Any, Dict, List, Optional, Tuple

# Third party imports
from openpyxl import load_workbook

# Module imports
from plane.db.models import Branch, Entity, Project, SteeringProfileRole, Workspace

from .people import ADMIN, MEMBER, PeopleError, PersonSpec, ProfileSpec, clean_email

MAX_ROWS = 1000

# Column headers, French first (the template), English accepted.
COLUMNS = {
    "first_name": ("prenom", "first name", "firstname"),
    "last_name": ("nom", "last name", "lastname"),
    "email": ("e-mail", "email", "adresse e-mail", "mail"),
    "role": ("role", "role dans l'organisation"),
    "profile": ("profil", "profil de pilotage", "profile"),
    "scope": ("perimetre", "pole ou entite", "scope"),
    "projects": ("projets", "projects"),
}
TEMPLATE_HEADER = ["Prénom", "Nom", "E-mail", "Rôle", "Profil", "Périmètre", "Projets"]

ROLE_LABELS = {"admin": ADMIN, "administrateur": ADMIN, "membre": MEMBER, "member": MEMBER, "": MEMBER}
PROFILE_LABELS = {
    "pdg": SteeringProfileRole.CEO,
    "ceo": SteeringProfileRole.CEO,
    "cabinet": SteeringProfileRole.CABINET,
    "cabinet de la presidence": SteeringProfileRole.CABINET,
    "directeur de pole": SteeringProfileRole.BRANCH_DIRECTOR,
    "directrice de pole": SteeringProfileRole.BRANCH_DIRECTOR,
    "branch_director": SteeringProfileRole.BRANCH_DIRECTOR,
    "responsable d'entite": SteeringProfileRole.ENTITY_MANAGER,
    "entity_manager": SteeringProfileRole.ENTITY_MANAGER,
}
ROLE_EXPORT = {ADMIN: "Admin", MEMBER: "Membre"}
STATUS_EXPORT = {"active": "Actif", "pending": "Jamais connecté", "departed": "Parti"}
PROFILE_EXPORT = {
    SteeringProfileRole.CEO: "PDG",
    SteeringProfileRole.CABINET: "Cabinet de la Présidence",
    SteeringProfileRole.BRANCH_DIRECTOR: "Directeur de pôle",
    SteeringProfileRole.ENTITY_MANAGER: "Responsable d'entité",
}


def normalize(value: Any) -> str:
    """Lower case, no accents, single spaces: « Pôle » and « pole » match."""
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return " ".join(text.replace("’", "'").lower().split())


def read_table(filename: str, content: bytes) -> List[List[str]]:
    """Rows of a .csv (comma or semicolon, UTF-8 or Windows encoding) or the first sheet of a .xlsx."""
    if filename.lower().endswith(".xlsx"):
        try:
            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        except Exception:
            raise PeopleError("The Excel file cannot be read.", "unreadable_file")
        return [
            ["" if cell is None else str(cell).strip() for cell in row] for row in sheet.iter_rows(values_only=True)
        ]
    if not filename.lower().endswith(".csv"):
        raise PeopleError("Use a .csv or .xlsx file.", "unsupported_file")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise PeopleError("The CSV file cannot be read.", "unreadable_file")
    delimiter = ";" if text.split("\n", 1)[0].count(";") > text.split("\n", 1)[0].count(",") else ","
    return [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _header_map(header: List[str]) -> Dict[str, int]:
    positions = {}
    for index, title in enumerate(header):
        name = normalize(title)
        for key, aliases in COLUMNS.items():
            if name in aliases and key not in positions:
                positions[key] = index
    missing = [key for key in ("first_name", "email") if key not in positions]
    if missing:
        raise PeopleError("The file needs at least the columns Prénom and E-mail.", "missing_columns")
    return positions


class Lookups:
    """Names of the referential and projects, matched without case or accents."""

    def __init__(self, workspace: Workspace):
        self.branches = {normalize(b.name): str(b.id) for b in Branch.objects.filter(workspace=workspace)}
        self.entities = {normalize(e.name): str(e.id) for e in Entity.objects.filter(workspace=workspace)}
        self.projects: Dict[str, str] = {}
        for project in Project.objects.filter(workspace=workspace, archived_at__isnull=True):
            self.projects[normalize(project.identifier)] = str(project.id)
            self.projects.setdefault(normalize(project.name), str(project.id))


def parse_rows(workspace: Workspace, table: List[List[str]]) -> List[Dict[str, Any]]:
    """
    One entry per data line: the values read, the resulting person (or None) and the errors, so
    the interface can show the whole file before anything is written.
    """
    rows = [row for row in table if any(cell for cell in row)]
    if not rows:
        raise PeopleError("The file is empty.", "empty_file")
    if len(rows) > MAX_ROWS + 1:
        raise PeopleError(f"At most {MAX_ROWS} people per import.", "too_many_rows")
    positions = _header_map(rows[0])
    lookups = Lookups(workspace)
    seen_emails = set()
    parsed = []
    for line, row in enumerate(rows[1:], start=2):

        def cell(key: str, row=row) -> str:
            index = positions.get(key)
            return row[index].strip() if index is not None and index < len(row) else ""

        values = {key: cell(key) for key in COLUMNS}
        spec, errors = _to_spec(values, lookups)
        if spec and spec.email in seen_emails:
            errors.append("duplicate_in_file")
        if spec:
            seen_emails.add(spec.email)
        parsed.append({"line": line, "values": values, "spec": spec if not errors else None, "errors": errors})
    return parsed


def _to_spec(values: Dict[str, str], lookups: Lookups) -> Tuple[Optional[PersonSpec], List[str]]:
    errors = []
    try:
        email = clean_email(values["email"])
    except PeopleError:
        email, errors = "", ["invalid_email"]
    if not values["first_name"]:
        errors.append("missing_first_name")
    role = ROLE_LABELS.get(normalize(values["role"]))
    if role is None:
        errors.append("invalid_role")

    profiles = []
    if values["profile"]:
        profile_role = PROFILE_LABELS.get(normalize(values["profile"]))
        scope = normalize(values["scope"])
        if not profile_role:
            errors.append("invalid_profile")
        elif profile_role == SteeringProfileRole.BRANCH_DIRECTOR:
            if scope in lookups.branches:
                profiles.append(ProfileSpec(role=profile_role, branch_id=lookups.branches[scope]))
            else:
                errors.append("unknown_branch")
        elif profile_role == SteeringProfileRole.ENTITY_MANAGER:
            if scope in lookups.entities:
                profiles.append(ProfileSpec(role=profile_role, entity_id=lookups.entities[scope]))
            else:
                errors.append("unknown_entity")
        else:
            profiles.append(ProfileSpec(role=profile_role))

    project_ids = []
    for name in values["projects"].replace(",", ";").split(";"):
        if not name.strip():
            continue
        project_id = lookups.projects.get(normalize(name))
        if project_id:
            project_ids.append(project_id)
        else:
            errors.append("unknown_project")
            break

    if errors:
        return None, errors
    return (
        PersonSpec(
            email=email,
            first_name=values["first_name"][:255],
            last_name=values["last_name"][:255],
            role=role,
            profiles=profiles,
            project_ids=list(dict.fromkeys(project_ids)),
        ),
        [],
    )


def export_csv(
    people: List[Dict[str, Any]], branches: Dict[str, str], entities: Dict[str, str], projects: Dict[str, str]
) -> str:
    """The people list in the import format (one profile per line: the first one), plus the status."""
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow([*TEMPLATE_HEADER, "Statut"])
    for person in people:
        profile = person["profiles"][0] if person["profiles"] else None
        scope = ""
        if profile and profile["branch_id"]:
            scope = branches.get(profile["branch_id"], "")
        if profile and profile["entity_id"]:
            scope = entities.get(profile["entity_id"], "")
        writer.writerow(
            [
                person["first_name"],
                person["last_name"],
                person["email"],
                ROLE_EXPORT.get(person["role"], ""),
                PROFILE_EXPORT.get(profile["role"], "") if profile else "",
                scope,
                "; ".join(projects.get(p["project_id"], "") for p in person["projects"]),
                STATUS_EXPORT.get(person["status"], person["status"]),
            ]
        )
    return "﻿" + output.getvalue()
